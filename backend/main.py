"""Dude AI - simple Jira connection and issue importer for learning."""
import os
import json
from contextlib import asynccontextmanager
from pathlib import Path
import sqlite3
import time
from datetime import datetime
from threading import Lock
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import httpx
from fastapi import FastAPI, HTTPException, Query, Request, Response
from fastapi.middleware.cors import CORSMiddleware
from starlette.responses import JSONResponse
from pydantic import BaseModel, HttpUrl, Field

DATA_DIR = Path(os.environ.get("DATA_DIR", Path(__file__).parent)).resolve()
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB_FILE = DATA_DIR / "dude_ai.db"
EXCEL_FILE = DATA_DIR / "api_audit_tracker.xlsx"
EXCEL_LOCK = Lock()
ACTIVE_ACTIONS = set()
ACTIVE_ACTIONS_LOCK = Lock()


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="174A45")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12), 52)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def ensure_excel_tracker():
    EXCEL_FILE.parent.mkdir(parents=True, exist_ok=True)
    if EXCEL_FILE.exists():
        # Migrate workbooks created by earlier versions without deleting existing audit history.
        with EXCEL_LOCK:
            wb = load_workbook(EXCEL_FILE)
            required = {
                "API Call Log": ["Timestamp", "User Name", "User ID", "Method", "Path", "Status", "Duration ms", "Frontend action", "Click ID", "Blocked", "Frontend duration ms", "Backend duration ms", "Jira duration ms", "Test case ID", "Environment", "Notes"],
                "Jira API Call Log": ["Timestamp", "Frontend action", "Jira endpoint", "HTTP status", "Duration ms", "Result", "Notes"],
                "Button Click Log": ["Timestamp", "User Name", "User ID", "Button / Action", "Click ID", "Click number", "Outcome", "Block reason", "Frontend route", "API endpoint", "HTTP status", "Duration ms", "Test run ID"],
                "Button Click Summary": ["Button / Action", "Total clicks", "Accepted clicks", "Blocked clicks", "Backend calls", "Jira calls", "Pass count", "Fail count", "Notes"],
                "QA Test Cases": ["Test Case ID", "Area", "Scenario", "Test Type", "Priority", "Preconditions", "Steps", "Expected Result", "Actual Result", "Status", "Evidence / Defect ID", "Standard / Rationale"],
                "Imported Jira Issues": ["Imported At", "Issue Key", "Project Key", "Summary", "Status", "Assignee", "Updated At", "Jira URL", "Imported By", "User ID"],
                "Frontend Backend Map": ["Frontend file", "Frontend function / action", "HTTP method", "API endpoint", "Backend handler", "Purpose"],
            }
            changed = False
            for name, headers in required.items():
                if name not in wb.sheetnames:
                    ws = wb.create_sheet(name)
                    ws.append(headers)
                    changed = True
                else:
                    ws = wb[name]
                    current_headers = [cell.value for cell in ws[1]]
                    if "User Name" in headers and "User Name" not in current_headers:
                        for col_idx, h_val in enumerate(headers, 1):
                            ws.cell(row=1, column=col_idx, value=h_val)
                        changed = True
                    elif "Imported By" in headers and "Imported By" not in current_headers:
                        for col_idx, h_val in enumerate(headers, 1):
                            ws.cell(row=1, column=col_idx, value=h_val)
                        changed = True
            if "API Call Log" in wb.sheetnames and wb["API Call Log"].max_row == 1:
                wb["API Call Log"][1][0].value = "Timestamp"
            if changed:
                for ws in wb.worksheets:
                    style_sheet(ws)
                wb.save(EXCEL_FILE)
        return
    wb = Workbook()
    api = wb.active
    api.title = "API Call Log"
    api.append(["Timestamp", "User Name", "User ID", "Method", "Path", "Status", "Duration ms", "Frontend action", "Click ID", "Blocked", "Frontend duration ms", "Backend duration ms", "Jira duration ms", "Test case ID", "Environment", "Notes"])
    jira_log = wb.create_sheet("Jira API Call Log")
    jira_log.append(["Timestamp", "Frontend action", "Jira endpoint", "HTTP status", "Duration ms", "Result", "Notes"])
    clicks = wb.create_sheet("Button Click Log")
    clicks.append(["Timestamp", "User Name", "User ID", "Button / Action", "Click ID", "Click number", "Outcome", "Block reason", "Frontend route", "API endpoint", "HTTP status", "Duration ms", "Test run ID"])
    summary = wb.create_sheet("Button Click Summary")
    summary.append(["Button / Action", "Total clicks", "Accepted clicks", "Blocked clicks", "Backend calls", "Jira calls", "Pass count", "Fail count", "Notes"])
    qa = wb.create_sheet("QA Test Cases")
    qa.append(["Test Case ID", "Area", "Scenario", "Test Type", "Priority", "Preconditions", "Steps", "Expected Result", "Actual Result", "Status", "Evidence / Defect ID", "Standard / Rationale"])
    qa_rows = [
        ["UI-001", "UI", "Primary button is visible and usable", "UI / Functional", "High", "App loaded", "Locate button by accessible role/name and click", "Action starts once; user receives clear feedback", "", "Not Run", "", "Playwright user-visible testing and actionability"],
        ["UI-002", "UI", "Button is protected during an in-flight request", "Negative / Concurrency", "Critical", "API response delayed in test", "Click the same action repeatedly while pending", "Only one action request starts; extra clicks are recorded as blocked", "", "Not Run", "", "Duplicate-action protection and concurrency testing"],
        ["API-001", "API", "Valid request returns expected response", "API / Positive", "High", "Valid test credentials and authorized test Jira", "Send request once and verify status and schema", "Expected 2xx response and required fields", "", "Not Run", "", "API contract and functional validation"],
        ["API-002", "API", "Invalid credentials are handled safely", "API / Negative", "High", "Test-only invalid credential", "Submit invalid credentials", "Clear error; secret not logged; no credential persistence", "", "Not Run", "", "Security and negative testing"],
        ["API-003", "API", "Concurrent duplicate API requests are rejected", "API / Concurrency", "Critical", "Authorized test environment", "Send concurrent identical action requests", "One request proceeds; duplicate gets 409 and is logged", "", "Not Run", "", "OWASP API resource-consumption/rate-limit testing principles"],
        ["API-004", "API", "Request timing and status are logged", "Observability", "High", "Excel tracker writable", "Run a known test request", "Timestamp, endpoint, status, action and duration recorded", "", "Not Run", "", "ISTQB test logging and execution timing"],
        ["DATA-001", "Data", "Imported issues are not duplicated on repeat import", "Data Integrity", "High", "Seeded test project", "Import the same issues twice", "Rows update by issue key rather than duplicating", "", "Not Run", "", "Idempotency and data-integrity testing"],
        ["SEC-001", "Security", "Credentials and tokens are excluded from logs", "Security / Privacy", "Critical", "Test credentials", "Perform connection test and inspect logs/workbook", "No API token, password or request body appears in logs", "", "Not Run", "", "Sensitive-data minimization"],
        ["E2E-001", "End-to-end", "Connect, load projects, import and view issues", "E2E / Smoke", "High", "Dedicated Jira test project", "Run the user journey", "Projects and imported issues display correctly; timings recorded", "", "Not Run", "", "Playwright isolated tests and user-visible assertions"],
        ["ERR-001", "Resilience", "Jira timeout or server error is handled", "Negative / Resilience", "Medium", "Mock or controlled failing endpoint", "Simulate timeout / 5xx", "Action finishes safely, error shown, lock released", "", "Not Run", "", "Error-path and recovery testing"],
        ["USR-001", "Identity", "Welcome username and user ID collected", "User Flow / Identity", "High", "App loaded", "Enter username and user ID and submit", "Personalized greeting shown; headers passed to backend; internal ID logged", "", "Not Run", "", "User identity onboarding and audit attribution"],
        ["VULN-001", "Security / Import", "Import vulnerability JSON into Jira project", "Integration / Import", "High", "Valid JSON and Jira connection", "Upload security_vulnerabilities.json, preview, and confirm import", "Issues created in Jira; duplicates handled; audit tracker updated", "", "Not Run", "", "Vulnerability import workflow and audit integrity"],
    ]
    for row in qa_rows:
        qa.append(row)
    docs = wb.create_sheet("Imported Jira Issues")
    docs.append(["Imported At", "Issue Key", "Project Key", "Summary", "Status", "Assignee", "Updated At", "Jira URL", "Imported By", "User ID"])
    mapping = wb.create_sheet("Frontend Backend Map")
    mapping.append(["Frontend file", "Frontend function / action", "HTTP method", "API endpoint", "Backend handler", "Purpose"])
    mapping_rows = [
        ["frontend/src/App.jsx", "Welcome / user identity flow", "HEADER", "X-User-Name, X-User-ID", "record_api_call()", "Pass user identity for audit attribution and Jira mapping."],
        ["frontend/src/App.jsx", "request() on Test connection", "POST", "/api/jira/test-connection", "test_connection()", "Verify Jira email and API token, validate entered user ID against Jira accountId."],
        ["frontend/src/App.jsx", "loadProjects()", "POST", "/api/jira/projects", "get_projects()", "List Jira projects the account can access."],
        ["frontend/src/App.jsx", "fetchDocuments()", "POST", "/api/jira/fetch-documents", "fetch_documents()", "Fetch Jira issues and save/update local records with user attribution."],
        ["frontend/src/App.jsx", "validateImportJson()", "POST", "/api/jira/validate-import-json", "validate_import_json()", "Validate uploaded Jira vulnerability JSON structure and fields."],
        ["frontend/src/App.jsx", "importVulnerabilities()", "POST", "/api/jira/import-vulnerabilities", "import_vulnerabilities()", "Create Jira issues individually with ADF descriptions and duplicate prevention."],
        ["frontend/src/App.jsx", "loadDocuments()", "GET", "/api/documents?limit=500", "documents()", "Load locally imported issues for the Documents page."],
        ["frontend/src/App.jsx", "topbar refresh button", "GET", "/api/documents?limit=500", "documents()", "Refresh the visible local issue list."],
        ["frontend/src/App.jsx", "Download Excel report button", "GET", "/api/reports/excel", "download_excel_report()", "Download the api_audit_tracker.xlsx workbook."],
        ["Browser / health check", "API status", "GET", "/api/health", "health()", "Check whether the local backend is running."],
        ["Browser / API docs", "Swagger UI", "GET", "/docs", "FastAPI generated docs", "Inspect and test endpoints locally."],
    ]
    for row in mapping_rows:
        mapping.append(row)
    for ws in wb.worksheets:
        style_sheet(ws)
    with EXCEL_LOCK:
        wb.save(EXCEL_FILE)

def append_excel(sheet_name, row):
    ensure_excel_tracker()
    with EXCEL_LOCK:
        wb = load_workbook(EXCEL_FILE)
        ws = wb[sheet_name]
        ws.append(row)
        style_sheet(ws)
        wb.save(EXCEL_FILE)

def log_api_call(method: str, path: str, status: int, duration_ms: float, frontend_action: str = "", click_id: str = "", blocked: bool = False, user_name: str = "", user_id: str = ""):
    append_excel("API Call Log", [datetime.now().astimezone().isoformat(timespec="seconds"), user_name, user_id, method, path, status, round(duration_ms, 2), frontend_action, click_id, "Yes" if blocked else "No", "", round(duration_ms, 2), "", "", "Cloud" if "onrender.com" in path else "Local", "No request body or credentials are recorded."])

def log_jira_api_call(action: str, endpoint: str, status: int, duration_ms: float):
    append_excel("Jira API Call Log", [datetime.now().astimezone().isoformat(timespec="seconds"), action, endpoint, status, round(duration_ms, 2), "Success" if 200 <= status < 400 else "Failed", "Endpoint only; credentials and request bodies are never logged."])


async def timed_jira_get(client: httpx.AsyncClient, endpoint: str, action: str, params=None):
    started = time.perf_counter()
    status = 0
    try:
        response = await client.get(endpoint, params=params)
        status = response.status_code
        return response
    finally:
        elapsed = (time.perf_counter() - started) * 1000
        try:
            log_jira_api_call(action, endpoint, status, elapsed)
        except Exception as exc:
            print(f"[Jira API audit log warning] {exc}")


async def timed_jira_post(client: httpx.AsyncClient, endpoint: str, action: str, json_data=None):
    started = time.perf_counter()
    status = 0
    try:
        response = await client.post(endpoint, json=json_data)
        status = response.status_code
        return response
    finally:
        elapsed = (time.perf_counter() - started) * 1000
        try:
            log_jira_api_call(action, endpoint, status, elapsed)
        except Exception as exc:
            print(f"[Jira API audit log warning] {exc}")



def log_button_event(event: dict):
    append_excel("Button Click Log", [
        datetime.now().astimezone().isoformat(timespec="seconds"),
        event.get("user_name", ""),
        event.get("user_id", ""),
        event.get("button_name", "Unknown"),
        event.get("click_id", ""),
        event.get("click_number", 1),
        event.get("outcome", "accepted"),
        event.get("block_reason", ""),
        event.get("frontend_route", ""),
        event.get("api_endpoint", ""),
        event.get("http_status", ""),
        event.get("duration_ms", ""),
        event.get("test_run_id", "")
    ])
    # Recalculate the summary from the detailed event log so blocked attempts are counted too.
    with EXCEL_LOCK:
        wb = load_workbook(EXCEL_FILE)
        ws = wb["Button Click Log"]
        header_row = [str(c.value or "").strip().lower() for c in ws[1]]
        name_col = header_row.index("button / action") if "button / action" in header_row else 3
        click_col = header_row.index("click id") if "click id" in header_row else 4
        outcome_col = header_row.index("outcome") if "outcome" in header_row else 6

        by_click = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) <= name_col:
                continue
            name = row[name_col] or "Unknown"
            click_id = row[click_col] if len(row) > click_col and row[click_col] else ""
            outcome = row[outcome_col] if len(row) > outcome_col and row[outcome_col] else "accepted"
            identity = click_id or f"{name}:{row[0]}"
            item = by_click.setdefault((name, identity), {"blocked": False, "accepted": False})
            if outcome == "blocked":
                item["blocked"] = True
            else:
                item["accepted"] = True
        counts = {}
        for (name, _identity), outcome_flags in by_click.items():
            item = counts.setdefault(name, {"total": 0, "accepted": 0, "blocked": 0})
            item["total"] += 1
            if outcome_flags["blocked"] and not outcome_flags["accepted"]:
                item["blocked"] += 1
            else:
                item["accepted"] += 1
        summary = wb["Button Click Summary"]
        if summary.max_row > 1:
            summary.delete_rows(2, summary.max_row - 1)
        for name, c in sorted(counts.items()):
            summary.append([name, c["total"], c["accepted"], c["blocked"], "See API Call Log", "See Jira endpoint trace", "", "", "Counts from Button Click Log"])
        style_sheet(summary)
        wb.save(EXCEL_FILE)

def log_imported_issues(records):
    ensure_excel_tracker()
    with EXCEL_LOCK:
        wb = load_workbook(EXCEL_FILE)
        ws = wb["Imported Jira Issues"]
        # Replace existing log rows for the issue keys in this fetch, so repeated imports do not duplicate them.
        issue_keys = {record["issue_key"] for record in records}
        for row_index in range(ws.max_row, 1, -1):
            if ws.cell(row_index, 2).value in issue_keys:
                ws.delete_rows(row_index, 1)
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        for record in records:
            ws.append([now, record["issue_key"], record["project_key"], record["summary"], record["status"], record["assignee"], record["updated_at"], record["browse_url"], record.get("imported_by", ""), record.get("imported_by_id", "")])
        style_sheet(ws)
        wb.save(EXCEL_FILE)


def db():
    con = sqlite3.connect(DB_FILE)
    con.row_factory = sqlite3.Row
    return con


def init_db():
    with db() as con:
        con.execute("""CREATE TABLE IF NOT EXISTS documents (
            issue_key TEXT PRIMARY KEY, project_key TEXT NOT NULL, summary TEXT NOT NULL,
            description TEXT DEFAULT '', status TEXT DEFAULT '', assignee TEXT DEFAULT '',
            updated_at TEXT DEFAULT '', browse_url TEXT DEFAULT '',
            imported_by TEXT DEFAULT '', imported_by_id TEXT DEFAULT '')""")
        # Safe migration if table already existed without user columns
        try:
            cursor = con.execute("PRAGMA table_info(documents)")
            cols = {row["name"] for row in cursor.fetchall()}
            if "imported_by" not in cols:
                con.execute("ALTER TABLE documents ADD COLUMN imported_by TEXT DEFAULT ''")
            if "imported_by_id" not in cols:
                con.execute("ALTER TABLE documents ADD COLUMN imported_by_id TEXT DEFAULT ''")
        except Exception:
            pass


init_db()


@asynccontextmanager
async def lifespan(app):
    init_db()
    try:
        ensure_excel_tracker()
        print(f"[Dude AI] API audit workbook: {EXCEL_FILE}")
    except Exception as exc:
        print(f"[Dude AI warning] Could not initialize Excel audit workbook: {exc}")
    yield

app = FastAPI(title="Dude AI API", description="Connect Jira and import project issues.", lifespan=lifespan)

@app.middleware("http")
async def record_api_call(request: Request, call_next):
    started = time.perf_counter()
    status_code = 500
    path = request.url.path + (f"?{request.url.query}" if request.url.query else "")
    action_key = request.headers.get("x-frontend-action", request.url.path)
    click_id = request.headers.get("x-click-id", "")
    user_name = request.headers.get("x-user-name", "")
    user_id = request.headers.get("x-user-id", "")
    guarded = request.method == "POST" and request.url.path in {"/api/jira/test-connection", "/api/jira/projects", "/api/jira/fetch-documents", "/api/jira/import-vulnerabilities"}
    guard_key = request.url.path
    acquired = False
    try:
        if guarded:
            with ACTIVE_ACTIONS_LOCK:
                if guard_key in ACTIVE_ACTIONS:
                    status_code = 409
                    return JSONResponse(
                        {"detail": "Duplicate action blocked: a request for this action is already running."},
                        status_code=409,
                    )
                ACTIVE_ACTIONS.add(guard_key)
                acquired = True
        response = await call_next(request)
        status_code = response.status_code
        return response
    finally:
        elapsed = (time.perf_counter() - started) * 1000
        if acquired:
            with ACTIVE_ACTIONS_LOCK:
                ACTIVE_ACTIONS.discard(guard_key)
        if path.startswith("/api/") and request.url.path != "/api/telemetry/button-event":
            try:
                log_api_call(request.method, path, status_code, elapsed, action_key, click_id, status_code == 409, user_name=user_name, user_id=user_id)
                # Record backend-side duplicate rejections in the click log as well,
                # including requests that bypass the React frontend guard.
                if guarded and status_code == 409:
                    log_button_event({
                        "user_name": user_name,
                        "user_id": user_id,
                        "button_name": action_key or request.url.path,
                        "click_id": click_id or f"backend-blocked-{time.time_ns()}",
                        "click_number": 0,
                        "outcome": "blocked",
                        "block_reason": "Backend rejected overlapping request (HTTP 409).",
                        "frontend_route": request.headers.get("referer", ""),
                        "api_endpoint": request.url.path,
                        "http_status": 409,
                        "duration_ms": round(elapsed, 2),
                        "test_run_id": request.headers.get("x-test-run-id", "")
                    })
            except Exception as exc:
                print(f"[Excel audit log warning] {exc}")

# Configure CORS origins for Vercel production, preview deployments, Render, and local dev
configured_origins = [
    "https://dude-ai-roan.vercel.app",
    "https://dudeai.onrender.com",
    "https://dude-ai.vercel.app",
    "https://dude-k6fe00oce-supporticon.vercel.app",
    "https://dude-ai-git-main-supporticon.vercel.app",
    "http://127.0.0.1:5173",
    "http://localhost:5173",
    "http://127.0.0.1:3000",
    "http://localhost:3000",
    "http://127.0.0.1:4173",
    "http://localhost:4173",
    "http://127.0.0.1:8000",
    "http://localhost:8000",
]

env_origins = os.environ.get("ALLOWED_ORIGINS") or os.environ.get("CORS_ORIGINS", "")
if env_origins:
    for origin in env_origins.split(","):
        cleaned = origin.strip().rstrip("/")
        if cleaned and cleaned not in configured_origins:
            configured_origins.append(cleaned)

# Added outermost so that all responses (including 409 rejections) contain CORS headers
app.add_middleware(
    CORSMiddleware,
    allow_origins=configured_origins,
    allow_origin_regex=r"^https:\/\/(dude-[a-zA-Z0-9_-]+|dudeai[a-zA-Z0-9_-]*)\.vercel\.app$",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["*"],
)



class ButtonEvent(BaseModel):
    user_name: str = ""
    user_id: str = ""
    button_name: str
    click_id: str = ""
    click_number: int = 1
    outcome: str = "accepted"
    block_reason: str = ""
    frontend_route: str = ""
    api_endpoint: str = ""
    http_status: int | str = ""
    duration_ms: float | str = ""
    test_run_id: str = ""


@app.post("/api/telemetry/button-event")
def button_event(body: ButtonEvent):
    log_button_event(body.model_dump())
    return {"success": True}


class Credentials(BaseModel):
    base_url: HttpUrl
    email: str = Field(min_length=3)
    api_token: str = Field(min_length=1)
    user_name: str | None = None
    user_id: str | None = None


class FetchRequest(Credentials):
    project_key: str = Field(min_length=1, max_length=30)
    max_results: int = Field(default=100, ge=1, le=100)


class SearchRequest(BaseModel):
    q: str = ""
    project_key: str | None = None


class ValidateJsonRequest(BaseModel):
    raw_json: str | None = None
    issues_data: list[dict] | dict | None = None


class ImportVulnerabilitiesRequest(Credentials):
    target_project_key: str = Field(min_length=1, max_length=30)
    issues: list[dict] = Field(min_length=1)
    skip_duplicates: bool = True


def clean_url(value: str) -> str:
    value = value.rstrip("/")
    parsed = urlparse(value)
    if not parsed.netloc or not parsed.scheme.startswith("http"):
        raise HTTPException(400, "Enter a valid Jira site URL, for example https://your-site.atlassian.net")
    return value


def jira_client(body: Credentials):
    return httpx.AsyncClient(base_url=clean_url(str(body.base_url)), auth=(body.email, body.api_token),
                             headers={"Accept": "application/json"}, timeout=30.0)


def jira_error(response: httpx.Response, action: str):
    if response.status_code == 401:
        detail = "Jira rejected the email or API token. Check both credentials."
    elif response.status_code == 403:
        detail = "Your Jira account does not have permission to perform this action."
    elif response.status_code == 410:
        detail = f"Jira returned HTTP 410 for {action}. Check that the current Jira REST endpoint is being used."
    else:
        detail = f"Jira {action} failed with HTTP {response.status_code}: {response.text[:300]}"
    raise HTTPException(response.status_code if response.status_code < 500 else 502, detail)


def to_text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return " ".join(filter(None, (to_text(item) for item in value)))
    if isinstance(value, dict):
        if value.get("type") == "text":
            return value.get("text", "")
        return " ".join(filter(None, (to_text(v) for k, v in value.items() if k not in {"type", "version"})))
    return str(value)


def text_to_adf(value: str) -> dict:
    text = (value or "").strip()
    if not text:
        text = "No description provided."
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not paragraphs:
        paragraphs = [text]
    return {
        "type": "doc",
        "version": 1,
        "content": [
            {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": p}
                ]
            }
            for p in paragraphs
        ]
    }



@app.get("/")
def home():
    return {"app": "Dude AI", "status": "running", "docs": "/docs"}


@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.get("/api/reports/excel")
def download_excel_report():
    try:
        ensure_excel_tracker()
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Failed to generate Excel audit tracker: {exc}")

    if not EXCEL_FILE.exists():
        raise HTTPException(status_code=404, detail="Excel audit tracker workbook not found.")

    with EXCEL_LOCK:
        try:
            with open(EXCEL_FILE, "rb") as f:
                content = f.read()
        except Exception as exc:
            raise HTTPException(status_code=500, detail=f"Failed to read Excel workbook: {exc}")

    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="api_audit_tracker.xlsx"',
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "Cache-Control": "no-cache, no-store, must-revalidate",
        },
    )


@app.post("/api/jira/test-connection")
async def test_connection(body: Credentials, request: Request):
    user_name = (body.user_name or request.headers.get("x-user-name", "")).strip()
    user_id = (body.user_id or request.headers.get("x-user-id", "")).strip()

    async with jira_client(body) as client:
        response = await timed_jira_get(client, "/rest/api/3/myself", "Test connection")
    if response.is_error:
        jira_error(response, "connection test")
    user = response.json()
    jira_display_name = user.get("displayName", body.email)
    jira_account_id = user.get("accountId", "")
    jira_email = user.get("emailAddress", body.email)

    identity_status = None
    if user_id or user_name:
        is_jira_account = False
        if user_id and user_id.lower() == jira_account_id.lower():
            is_jira_account = True
            msg = f"User ID '{user_id}' verified as your Jira Cloud account ID."
            status_tag = "verified_jira_account_id"
        elif user_id and user_id.lower() in [jira_email.lower(), jira_display_name.lower()]:
            msg = f"User ID matches Jira user '{jira_display_name}', but note Jira Cloud uses accountId ({jira_account_id[:8]}...) for issue assignments."
            status_tag = "matched_jira_profile"
        elif user_id:
            msg = f"User ID '{user_id}' is recorded as an internal application ID. It does not match the authenticated Jira account ID ({jira_account_id[:8]}...)."
            status_tag = "internal_application_id"
        else:
            msg = f"Active session for user '{user_name}'."
            status_tag = "session_user"

        identity_status = {
            "user_name": user_name,
            "user_id": user_id,
            "status": status_tag,
            "verified_jira_account": is_jira_account,
            "jira_account_id": jira_account_id,
            "jira_display_name": jira_display_name,
            "message": msg
        }

    return {
        "success": True,
        "display_name": jira_display_name,
        "jira_account_id": jira_account_id,
        "identity_status": identity_status,
        "message": "Connected successfully"
    }


@app.post("/api/jira/projects")
async def get_projects(body: Credentials):
    projects, start_at = [], 0
    async with jira_client(body) as client:
        while True:
            response = await timed_jira_get(client, "/rest/api/3/project/search", "Load projects", params={"startAt": start_at, "maxResults": 50, "orderBy": "name"})
            if response.is_error:
                jira_error(response, "project list")
            data = response.json()
            values = data.get("values", [])
            projects.extend({"id": str(p.get("id", "")), "key": p.get("key", ""), "name": p.get("name", ""), "project_type": p.get("projectTypeKey", "") } for p in values)
            start_at += len(values)
            if not values or start_at >= data.get("total", 0):
                break
    return {"success": True, "count": len(projects), "projects": projects}


@app.post("/api/jira/fetch-documents")
async def fetch_documents(body: FetchRequest, request: Request):
    user_name = (body.user_name or request.headers.get("x-user-name", "")).strip()
    user_id = (body.user_id or request.headers.get("x-user-id", "")).strip()
    project_key = body.project_key.strip().upper()
    jql = f'project = "{project_key}" ORDER BY updated DESC'
    fields = "summary,description,status,assignee,project,updated"
    issues, next_page_token = [], None
    async with jira_client(body) as client:
        # Jira Cloud's current enhanced JQL search endpoint is /search/jql.
        while len(issues) < body.max_results:
            params = {"jql": jql, "maxResults": min(100, body.max_results - len(issues)), "fields": fields}
            if next_page_token:
                params["nextPageToken"] = next_page_token
            response = await timed_jira_get(client, "/rest/api/3/search/jql", "Fetch documents", params=params)
            if response.is_error:
                jira_error(response, "issue search")
            data = response.json()
            batch = data.get("issues", [])
            issues.extend(batch)
            next_page_token = data.get("nextPageToken")
            if data.get("isLast", True) or not next_page_token or not batch:
                break
    records = []
    for issue in issues:
        fields_data = issue.get("fields") or {}
        status = fields_data.get("status") or {}
        assignee = fields_data.get("assignee") or {}
        records.append({
            "issue_key": issue.get("key", ""), "project_key": project_key,
            "summary": fields_data.get("summary") or "(No summary)",
            "description": to_text(fields_data.get("description")),
            "status": status.get("name", ""), "assignee": assignee.get("displayName", "Unassigned"),
            "updated_at": fields_data.get("updated", ""),
            "browse_url": f"{clean_url(str(body.base_url))}/browse/{issue.get('key', '')}",
            "imported_by": user_name,
            "imported_by_id": user_id,
        })
    with db() as con:
        con.executemany("""INSERT INTO documents(issue_key, project_key, summary, description, status, assignee, updated_at, browse_url, imported_by, imported_by_id)
            VALUES(:issue_key,:project_key,:summary,:description,:status,:assignee,:updated_at,:browse_url,:imported_by,:imported_by_id)
            ON CONFLICT(issue_key) DO UPDATE SET project_key=excluded.project_key, summary=excluded.summary,
            description=excluded.description, status=excluded.status, assignee=excluded.assignee,
            updated_at=excluded.updated_at, browse_url=excluded.browse_url,
            imported_by=excluded.imported_by, imported_by_id=excluded.imported_by_id""", records)
    log_imported_issues(records)
    print("\n" + "=" * 100)
    print(f"[Dude AI] Imported {len(records)} Jira issue(s) from project {project_key}")
    print(f"{'ISSUE KEY':<14} {'STATUS':<18} {'ASSIGNEE':<24} SUMMARY")
    print("-" * 100)
    for record in records:
        print(f"{record['issue_key']:<14} {record['status'][:17]:<18} {record['assignee'][:23]:<24} {record['summary']}")
    print(f"Excel tracker updated: {EXCEL_FILE}")
    print("=" * 100 + "\n")
    return {"success": True, "project_key": project_key, "fetched": len(records), "message": f"Imported {len(records)} issue(s)."}


@app.post("/api/jira/validate-import-json")
def validate_import_json(body: ValidateJsonRequest):
    data = None
    if body.issues_data is not None:
        data = body.issues_data
    elif body.raw_json:
        try:
            data = json.loads(body.raw_json)
        except Exception as exc:
            return {
                "valid": False,
                "count": 0,
                "issues": [],
                "detected_project_keys": [],
                "error": f"Invalid JSON syntax: {exc}",
                "message": "The provided content is not valid JSON."
            }
    else:
        return {
            "valid": False,
            "count": 0,
            "issues": [],
            "detected_project_keys": [],
            "error": "No JSON content or issue data provided.",
            "message": "Please select a JSON file or provide JSON content."
        }

    raw_list = []
    if isinstance(data, list):
        raw_list = data
    elif isinstance(data, dict):
        if "issues" in data and isinstance(data["issues"], list):
            raw_list = data["issues"]
        elif "fields" in data:
            raw_list = [data]
        else:
            raw_list = [v for v in data.values() if isinstance(v, dict)]

    if not raw_list:
        return {
            "valid": False,
            "count": 0,
            "issues": [],
            "detected_project_keys": [],
            "error": "No issue records found in JSON. Expected an array of issues or an object with an 'issues' property.",
            "message": "No issues found in the JSON file."
        }

    normalized = []
    detected_keys = set()
    for idx, item in enumerate(raw_list):
        if not isinstance(item, dict):
            continue
        fields = item.get("fields", item)
        summary = str(fields.get("summary") or "").strip()
        if not summary:
            continue

        orig_proj = ""
        proj_val = fields.get("project")
        if isinstance(proj_val, dict):
            orig_proj = str(proj_val.get("key", "")).strip().upper()
        elif isinstance(proj_val, str):
            orig_proj = proj_val.strip().upper()
        if orig_proj:
            detected_keys.add(orig_proj)

        issue_type = "Bug"
        it_val = fields.get("issuetype")
        if isinstance(it_val, dict):
            issue_type = it_val.get("name", "Bug")
        elif isinstance(it_val, str) and it_val.strip():
            issue_type = it_val.strip()

        priority = "High"
        p_val = fields.get("priority")
        if isinstance(p_val, dict):
            priority = p_val.get("name", "High")
        elif isinstance(p_val, str) and p_val.strip():
            priority = p_val.strip()

        labels = fields.get("labels", [])
        if isinstance(labels, str):
            labels = [l.strip() for l in labels.split(",") if l.strip()]
        elif not isinstance(labels, list):
            labels = []

        desc_raw = fields.get("description", "")
        desc_text = to_text(desc_raw)

        normalized.append({
            "index": idx + 1,
            "summary": summary,
            "issuetype": issue_type,
            "priority": priority,
            "original_project_key": orig_proj or "DAST",
            "labels": labels,
            "description": desc_text,
            "description_preview": desc_text[:140] + ("..." if len(desc_text) > 140 else ""),
            "raw": item
        })

    if not normalized:
        return {
            "valid": False,
            "count": 0,
            "issues": [],
            "detected_project_keys": [],
            "error": "No valid issues with a 'summary' field could be parsed from the JSON.",
            "message": "Zero valid issues found in JSON."
        }

    return {
        "valid": True,
        "count": len(normalized),
        "issues": normalized,
        "detected_project_keys": sorted(list(detected_keys)),
        "message": f"Successfully parsed and validated {len(normalized)} vulnerability issue(s)."
    }


@app.post("/api/jira/import-vulnerabilities")
async def import_vulnerabilities(body: ImportVulnerabilitiesRequest, request: Request):
    user_name = (body.user_name or request.headers.get("x-user-name", "")).strip()
    user_id = (body.user_id or request.headers.get("x-user-id", "")).strip()
    target_project_key = body.target_project_key.strip().upper()

    results = []
    created_records = []

    async with jira_client(body) as client:
        for item in body.issues:
            fields_in = item.get("fields", item)
            summary = str(fields_in.get("summary", "")).strip()
            if not summary:
                results.append({
                    "summary": "(Untitled issue)",
                    "status": "failed",
                    "error": "Issue summary is missing"
                })
                continue

            # Duplicate check: check local SQLite documents table
            if body.skip_duplicates:
                with db() as con:
                    cur = con.execute(
                        "SELECT issue_key FROM documents WHERE project_key = ? AND LOWER(summary) = LOWER(?) LIMIT 1",
                        (target_project_key, summary)
                    )
                    existing = cur.fetchone()
                    if existing:
                        results.append({
                            "summary": summary,
                            "status": "skipped",
                            "issue_key": existing["issue_key"],
                            "browse_url": f"{clean_url(str(body.base_url))}/browse/{existing['issue_key']}",
                            "message": f"Issue already exists in local records ({existing['issue_key']})"
                        })
                        continue

            # Map description to Jira ADF format
            desc_raw = fields_in.get("description", "")
            if isinstance(desc_raw, dict) and desc_raw.get("type") == "doc":
                adf_desc = desc_raw
            else:
                adf_desc = text_to_adf(to_text(desc_raw))

            issue_type_name = "Bug"
            if isinstance(fields_in.get("issuetype"), dict):
                issue_type_name = fields_in["issuetype"].get("name", "Bug")
            elif isinstance(fields_in.get("issuetype"), str) and fields_in["issuetype"].strip():
                issue_type_name = fields_in["issuetype"].strip()

            priority_name = "High"
            if isinstance(fields_in.get("priority"), dict):
                priority_name = fields_in["priority"].get("name", "High")
            elif isinstance(fields_in.get("priority"), str) and fields_in["priority"].strip():
                priority_name = fields_in["priority"].strip()

            labels = fields_in.get("labels", [])
            if isinstance(labels, str):
                labels = [l.strip() for l in labels.split(",") if l.strip()]

            jira_payload = {
                "fields": {
                    "project": {"key": target_project_key},
                    "summary": summary,
                    "issuetype": {"name": issue_type_name},
                    "description": adf_desc,
                }
            }
            if priority_name:
                jira_payload["fields"]["priority"] = {"name": priority_name}
            if labels:
                jira_payload["fields"]["labels"] = labels

            # Post individual issue to Jira Cloud API
            response = await timed_jira_post(client, "/rest/api/3/issue", "Create vulnerability issue", json_data=jira_payload)

            # Retry fallbacks if project issue types or priority are restricted
            if response.status_code == 400:
                err_text = response.text.lower()
                if "issuetype" in err_text and issue_type_name != "Task":
                    jira_payload["fields"]["issuetype"] = {"name": "Task"}
                    response = await timed_jira_post(client, "/rest/api/3/issue", "Create vulnerability issue (fallback Task)", json_data=jira_payload)
                if response.status_code == 400 and "priority" in response.text.lower():
                    jira_payload["fields"].pop("priority", None)
                    response = await timed_jira_post(client, "/rest/api/3/issue", "Create vulnerability issue (fallback no priority)", json_data=jira_payload)

            if response.is_error:
                err_detail = response.text[:250]
                try:
                    ej = response.json()
                    msgs = ej.get("errorMessages", [])
                    flds = [f"{k}: {v}" for k, v in ej.get("errors", {}).items()]
                    err_detail = " | ".join(filter(None, msgs + flds)) or err_detail
                except Exception:
                    pass
                results.append({
                    "summary": summary,
                    "status": "failed",
                    "error": f"Jira {response.status_code}: {err_detail}"
                })
            else:
                resp_json = response.json()
                created_key = resp_json.get("key", "")
                browse_url = f"{clean_url(str(body.base_url))}/browse/{created_key}"
                results.append({
                    "summary": summary,
                    "status": "created",
                    "issue_key": created_key,
                    "browse_url": browse_url,
                    "message": f"Successfully created {created_key}"
                })
                created_records.append({
                    "issue_key": created_key,
                    "project_key": target_project_key,
                    "summary": summary,
                    "description": to_text(fields_in.get("description", "")),
                    "status": "To Do",
                    "assignee": "Unassigned",
                    "updated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
                    "browse_url": browse_url,
                    "imported_by": user_name,
                    "imported_by_id": user_id,
                })

    if created_records:
        with db() as con:
            con.executemany("""INSERT INTO documents(issue_key, project_key, summary, description, status, assignee, updated_at, browse_url, imported_by, imported_by_id)
                VALUES(:issue_key,:project_key,:summary,:description,:status,:assignee,:updated_at,:browse_url,:imported_by,:imported_by_id)
                ON CONFLICT(issue_key) DO UPDATE SET project_key=excluded.project_key, summary=excluded.summary,
                description=excluded.description, status=excluded.status, assignee=excluded.assignee,
                updated_at=excluded.updated_at, browse_url=excluded.browse_url,
                imported_by=excluded.imported_by, imported_by_id=excluded.imported_by_id""", created_records)
        log_imported_issues(created_records)

    created_cnt = sum(1 for r in results if r["status"] == "created")
    skipped_cnt = sum(1 for r in results if r["status"] == "skipped")
    failed_cnt = sum(1 for r in results if r["status"] == "failed")

    return {
        "success": True,
        "target_project_key": target_project_key,
        "total": len(body.issues),
        "created": created_cnt,
        "skipped": skipped_cnt,
        "failed": failed_cnt,
        "results": results,
        "message": f"Import completed: {created_cnt} created, {skipped_cnt} skipped, {failed_cnt} failed."
    }



@app.get("/api/documents")
def documents(project_key: str | None = None, q: str = Query(default="", max_length=200), limit: int = Query(default=500, ge=1, le=500)):
    sql = "SELECT * FROM documents WHERE 1=1"
    args = []
    if project_key:
        sql += " AND project_key = ?"
        args.append(project_key.upper())
    if q.strip():
        sql += " AND (issue_key LIKE ? OR summary LIKE ? OR description LIKE ? OR status LIKE ?)"
        args.extend([f"%{q}%"] * 4)
    sql += " ORDER BY updated_at DESC LIMIT ?"
    args.append(limit)
    with db() as con:
        rows = [dict(row) for row in con.execute(sql, args).fetchall()]
    return {"count": len(rows), "documents": rows}


@app.get("/api/stats")
def stats():
    with db() as con:
        count = con.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
        projects = con.execute("SELECT COUNT(DISTINCT project_key) FROM documents").fetchone()[0]
    return {"documents": count, "projects": projects}


@app.delete("/api/documents/{project_key}")
def delete_project(project_key: str):
    with db() as con:
        cursor = con.execute("DELETE FROM documents WHERE project_key = ?", (project_key.upper(),))
    return {"success": True, "deleted": cursor.rowcount, "project_key": project_key.upper()}


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run("main:app", host="0.0.0.0", port=port, reload=False)
